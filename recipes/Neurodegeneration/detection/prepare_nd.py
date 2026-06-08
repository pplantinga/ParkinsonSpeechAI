import json
import os
import glob
import datetime
import numpy as np
import pathlib
import openpyxl
import pandas
import soundfile as sf
import collections

def summarize_split(json_dict, split_name):
    """Print per-dataset, per-class utterance and chunk counts for a split.

    'utts' counts unique utterance ids (key minus the trailing _<chunk>),
    matching how combine_chunks aggregates — so this is the count AUC and
    combined F-score are actually computed over. 'chunks' is the raw number
    of training/eval items.
    """
    chunk_counts = collections.Counter()
    utt_ids = collections.defaultdict(set)
    for key, entry in json_dict.items():
        info = entry["info_dict"]
        ds = info.get("dataset")
        pt = info.get("ptype")
        chunk_counts[(ds, pt)] += 1
        utt_ids[(ds, pt)].add(key.rsplit("_", 1)[0])

    print(f"\n[{split_name}] label distribution")
    print(f"  {'dataset':<10} {'class':<9} {'utts':>6} {'chunks':>8}")
    datasets = sorted({ds for ds, _ in chunk_counts}, key=lambda x: (x is None, str(x)))
    tot_utt = collections.Counter()
    for ds in datasets:
        for pt in ("Control", "Disease"):
            n_utt = len(utt_ids.get((ds, pt), ()))
            n_chunk = chunk_counts.get((ds, pt), 0)
            if n_utt or n_chunk:
                print(f"  {str(ds):<10} {str(pt):<9} {n_utt:>6} {n_chunk:>8}")
                tot_utt[pt] += n_utt
    print(f"  {'TOTAL':<10} {'Control':<9} {tot_utt['Control']:>6}")
    print(f"  {'TOTAL':<10} {'Disease':<9} {tot_utt['Disease']:>6}")

def convert_to_python(obj):
    """Recursively convert NumPy types to Python types for JSON serialization."""
    if isinstance(obj, np.integer):
        return int(obj)
    elif isinstance(obj, np.floating):
        return float(obj)
    elif isinstance(obj, np.ndarray):
        return obj.tolist()
    elif isinstance(obj, dict):
        return {k: convert_to_python(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [convert_to_python(item) for item in obj]
    else:
        return obj


def _norm_visit(v):
    """Normalize a visit / recording number to a canonical string.

    Handles ints (2), floats from Excel (2.0), and strings ("2") uniformly
    so that filename-derived and Excel-derived visit numbers compare equal.
    """
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _visit_sort_key(v):
    """Sort visits numerically when possible, lexically otherwise."""
    try:
        return (0, int(v))
    except (ValueError, TypeError):
        return (1, str(v))

def _encode_delaware_language(val):
    """Delaware PRIMARY LANGUAGE: 0 -> English, 1/2/other -> Other.
    A genuinely empty cell (None) stays None so 'no data' is distinct
    from 'non-English'."""
    if val is None:
        return None
    try:
        code = int(float(val))
    except (ValueError, TypeError):
        return "Other"
    return "English" if code == 0 else "Other"


def _encode_delaware_sex(val):
    """Delaware SEX: 1 -> Female, 2 -> Male, anything else -> None."""
    if val is None:
        return None
    try:
        code = int(float(val))
    except (ValueError, TypeError):
        return None
    return {1: "F", 2: "M"}.get(code)


# ---------------------------------------------------------------------------
# Top-level entry point
# ---------------------------------------------------------------------------

def prepare_nd(
    qpn_data_path,
    pitt_data_path,
    delaware_data_path,
    train_annotation,
    test_annotation,
    valid_annotation,
    chunk_size,
    delaware_demographics_path=None,
    delaware_split_seed=42,
    delaware_n_valid=8,
    delaware_n_test=8,
):
    """
    Build combined train/valid/test JSON manifests from the QPN, Pitt and
    Delaware datasets. Entries from all corpora are merged into a single
    JSON per split. Utterance keys are based on the source file path so
    they are unique across datasets without further prefixing.

    :param qpn_data_path: path to the QPN dataset root
    :param pitt_data_path: path to the Pitt dataset root
    :param delaware_data_path: path to the Delaware dataset root (must
        contain `MCI/` and `Control/` subdirectories of .wav files named
        like `<patient_id>-<recording_number>.wav`)
    :param train_annotation: output path for combined train JSON
    :param test_annotation: output path for combined test JSON
    :param valid_annotation: output path for combined valid JSON
    :param chunk_size: chunk size in seconds
    :param delaware_demographics_path: path to the Delaware demographics
        Excel (two sheets: Control then MCI). If None, a single *.xlsx in
        `delaware_data_path` is auto-detected.
    :param delaware_split_seed: RNG seed for the Delaware patient-level
        train/valid/test split. Same seed + same patient set => same split.
    :param delaware_n_valid: total number of patients in the valid split
        (balanced 50/50 across sex and ptype, so must be divisible by 4).
    :param delaware_n_test: total number of patients in the test split
        (balanced 50/50 across sex and ptype, so must be divisible by 4).
    """
    assert os.path.exists(qpn_data_path), "QPN data folder not found"
    assert os.path.exists(pitt_data_path), "Pitt data folder not found"
    assert os.path.exists(delaware_data_path), "Delaware data folder not found"

    # ---- QPN: gather per-split path -> info_dict ---------------------------
    qpn_path_type_dict = get_qpn_path_type_dicts(qpn_data_path)

    # ---- Pitt: gather per-split ground-truth dataframes --------------------
    pitt_data_path = pathlib.Path(pitt_data_path)
    pitt_csv = read_pitt_csv(pitt_data_path, "pitt_corpus")

    pitt_valid_ids = [59, 167, 211, 124, 182, 52, 208, 304,
                      122, 173, 238, 687, 10, 508, 213, 244]

    pitt_test_gt = pitt_csv[pitt_csv["test"] == 1]
    pitt_valid_gt = pitt_csv[
        (pitt_csv["test"] == 0) & (pitt_csv["id"].isin(pitt_valid_ids))
    ]
    pitt_train_gt = pitt_csv[
        (pitt_csv["test"] == 0) & (~pitt_csv["id"].isin(pitt_valid_ids))
    ]

    # ---- Delaware: load demographics, then stratified patient-level split --
    delaware_demographics = _load_delaware_demographics_or_none(
        delaware_data_path, delaware_demographics_path
    )
    delaware_splits = get_delaware_splits(
        delaware_data_path,
        demographics=delaware_demographics,
        n_valid=delaware_n_valid,
        n_test=delaware_n_test,
        seed=delaware_split_seed,
    )

    # ---- Build each combined split JSON ------------------------------------
    train_dict = create_combined_json(
        train_annotation,
        qpn_path_type_dict["train"],
        pitt_train_gt,
        delaware_splits["train"],
        chunk_size,
        overlap=None,
        pitt_overlap=None,
        delaware_overlap=None,
    )
    test_dict = create_combined_json(
        test_annotation,
        qpn_path_type_dict["test"],
        pitt_test_gt,
        delaware_splits["test"],
        chunk_size,
        overlap=None,
        pitt_overlap=None,
        delaware_overlap=None,
    )
    valid_dict = create_combined_json(
        valid_annotation,
        qpn_path_type_dict["valid"],
        pitt_valid_gt,
        delaware_splits["valid"],
        chunk_size,
        overlap=0,
        pitt_overlap=0,
        delaware_overlap=0,
    )

    summarize_split(train_dict, "train")
    summarize_split(valid_dict, "valid")
    summarize_split(test_dict, "test")

# ---------------------------------------------------------------------------
# QPN-specific helpers
# ---------------------------------------------------------------------------

def get_qpn_path_type_dicts(data_folder):
    """
    Extract the patient_type and other traits for each QPN recording.
    Returns a dict keyed by split name ('train'/'test'/'valid'), each
    mapping audio path -> patient trait dict.
    """
    datasets = os.listdir(data_folder)
    batch1_excel_path = os.path.join(data_folder, "QPN_Batch1.xlsx")
    batch2_excel_path = os.path.join(data_folder, "QPN_Batch2.xlsx")
    batch3_excel_path = os.path.join(data_folder, "QPN_Batch3.xlsx")
    path_type_dict = {}

    batch1_workbook = openpyxl.load_workbook(batch1_excel_path)
    batch1_sheet = batch1_workbook.active
    batch2_workbook = openpyxl.load_workbook(batch2_excel_path)
    batch2_sheet = batch2_workbook["Demographic"]
    batch3_workbook = openpyxl.load_workbook(batch3_excel_path)
    batch3_sheet = batch3_workbook.active

    for dataset in datasets:
        dataset_path = os.path.join(data_folder, dataset)

        if os.path.isfile(dataset_path):
            continue
        if dataset == "noise" or dataset == "rir":
            continue

        batch1_data_path = os.path.join(dataset_path, "Batch1")
        batch2_data_path = os.path.join(dataset_path, "Batch2")
        batch3_data_path = os.path.join(dataset_path, "Batch3")

        batch1_files = glob.glob(batch1_data_path + "/*.wav")
        batch2_files = glob.glob(batch2_data_path + "/*.wav")
        batch3_files = glob.glob(batch3_data_path + "/*.wav")

        batch1_patients = get_qpn_patient_traits(batch1_files, batch1_sheet, "Batch1")
        batch2_patients = get_qpn_patient_traits(batch2_files, batch2_sheet, "Batch2")
        batch3_patients = get_qpn_patient_traits(batch3_files, batch3_sheet, "Batch3")

        path_type_dict[dataset] = batch1_patients | batch2_patients | batch3_patients

    return path_type_dict


def get_qpn_patient_traits(files, sheet, batch):
    pids = [path.split("/")[-1].split("_")[1] for path in files]
    patients = {}

    for row in range(2, sheet.max_row + 1):  # Skip header row
        pid = sheet.cell(row=row, column=1).value
        ptype = sheet.cell(row=row, column=2).value
        sex = sheet.cell(row=row, column=3).value
        l1 = sheet.cell(row=row, column=4).value
        age = sheet.cell(row=row, column=6).value

        if pid is not None and pid.rstrip() in pids:
            ptype = ptype.rstrip()
            sex = sex.rstrip()
            l1 = l1.rstrip()

            # Refactor patient type
            if ptype == "CTRL" or ptype == "control":
                ptype = "Control"
            elif ptype == "PD" or ptype == "patient":
                ptype = "Disease"
            else:
                print(f"Unknown key found: {ptype}")
                continue

            # Refactor language
            if l1 == "FR" or "French" in l1 or "Fench" in l1:
                l1 = "French"
            elif l1 == "EN" or "English" in l1:
                l1 = "English"
            else:
                l1 = "Other"

            patients[pid] = {
                "ptype": ptype,
                "sex": sex,
                "age": age,
                "l1": l1,
            }

    # Map pids to file paths
    updated_dict = {}
    for pid in patients:
        for path in files:
            if pid in path:
                updated_dict[path] = patients[pid]

    return updated_dict


# ---------------------------------------------------------------------------
# Pitt-specific helpers
# ---------------------------------------------------------------------------

def read_pitt_csv(data_folder, subset):
    df = pandas.read_csv(data_folder / (subset + ".csv"))

    expanded_rows = []
    for _, row in df.iterrows():
        subfolder = "control" if row["dx"] == "Control" else "dementia"
        base_path = data_folder / subfolder
        id_str = str(row["id"]).zfill(3)
        pattern = f"{id_str}-*.wav"
        recording_files = sorted(glob.glob(str(base_path / pattern)))

        # For test set, only keep the first recording (-0)
        if row["test"] == 1:
            recording_files = [f for f in recording_files if f.endswith("-0.wav")]

        for rec_path in recording_files:
            new_row = row.copy()
            new_row["path"] = pathlib.Path(rec_path)
            expanded_rows.append(new_row)

    return pandas.DataFrame(expanded_rows)


# ---------------------------------------------------------------------------
# Delaware-specific helpers
# ---------------------------------------------------------------------------

# Header substring (lower-cased) -> unified info field. Matched against the
# header row so column order in the Excel doesn't matter. More specific
# substrings are listed first so they win ties.
DELAWARE_COLUMN_SPECS = {
    "record id": "pid",
    "visit number": "visit",
    "test date": "test_date",
    "age at testing": "age",
    "primary language": "l1",
    "moca": "moca",
    "sex": "sex",
}


def _load_delaware_demographics_or_none(delaware_data_path, demographics_path):
    """Resolve the demographics Excel path (explicit or auto-detected) and load it."""
    if demographics_path is None:
        candidates = glob.glob(os.path.join(delaware_data_path, "*.xlsx"))
        if len(candidates) == 1:
            demographics_path = candidates[0]
        elif len(candidates) > 1:
            print("[delaware] multiple .xlsx files found; pass "
                  f"delaware_demographics_path explicitly: {candidates}")
            return None
        else:
            print("[delaware] no demographics .xlsx found; proceeding without it")
            return None

    if not os.path.exists(demographics_path):
        print(f"[delaware] demographics file not found: {demographics_path}")
        return None

    return load_delaware_demographics(demographics_path)


def load_delaware_demographics(excel_path):
    """
    Load per-visit demographics from the Delaware Excel workbook.

    The workbook has two sheets: the FIRST holds Control participants and
    the SECOND holds MCI participants. A given RECORD ID can appear in BOTH
    classes, so demographics are stored per class and never matched across
    the class boundary.

    :param excel_path: path to the demographics .xlsx
    :return: {class_subfolder: {(pid, visit): info_dict}} where
        class_subfolder is "Control" or "MCI".
    """
    workbook = openpyxl.load_workbook(excel_path, data_only=True)

    # First sheet = Control, second sheet = MCI (per the workbook layout).
    sheet_by_class = {
        "Control": workbook.worksheets[0],
        "MCI": workbook.worksheets[1],
    }

    demographics = {}
    for class_name, sheet in sheet_by_class.items():
        # Build {field: column_index} from the header row.
        field_to_col = {}
        for col in range(1, sheet.max_column + 1):
            header = sheet.cell(row=1, column=col).value
            if header is None:
                continue
            header_norm = str(header).strip().lower()
            for substr, field in DELAWARE_COLUMN_SPECS.items():
                if substr in header_norm and field not in field_to_col:
                    field_to_col[field] = col
                    break

        for required in ("pid", "visit"):
            if required not in field_to_col:
                raise ValueError(
                    f"Could not find a '{required}' column in the "
                    f"'{class_name}' sheet of {excel_path}"
                )

        class_demo = {}
        for row in range(2, sheet.max_row + 1):
            record_id = sheet.cell(row=row, column=field_to_col["pid"]).value
            visit = sheet.cell(row=row, column=field_to_col["visit"]).value
            if record_id is None or visit is None:
                continue

            pid = _norm_visit(record_id) if isinstance(record_id, float) else str(record_id).strip()
            visit_str = _norm_visit(visit)

            info = {}
            for field, col in field_to_col.items():
                if field in ("pid", "visit"):
                    continue
                val = sheet.cell(row=row, column=col).value
                if field == "l1":
                    val = _encode_delaware_language(val)
                elif field == "sex":
                    val = _encode_delaware_sex(val)
                elif isinstance(val, (datetime.datetime, datetime.date)):
                    val = val.isoformat()      # keep JSON-serializable
                elif isinstance(val, str):
                    val = val.strip()
                info[field] = val
            info["visit"] = visit_str

            class_demo[(pid, visit_str)] = info

        demographics[class_name] = class_demo

    return demographics


def get_delaware_splits(
    data_folder, demographics=None, n_valid=8, n_test=8, seed=42
):
    """
    Walk the Delaware folder layout (MCI/<pid>-<rec>.wav,
    Control/<pid>-<rec>.wav), keep the EARLIEST available recording per
    patient, attach per-visit demographics, and bucket patients into
    train/valid/test with a deterministic sex-and-ptype-balanced split.

    Exactly n_valid patients go to valid and n_test patients go to test,
    balanced 50/50 across both sex (M/F) and ptype (MCI/Control). Each
    (ptype, sex) cell gets n_valid // 4 and n_test // 4 patients. Patients
    with unknown sex or that exceed the balanced quota go to train.

    :param data_folder: path to the Delaware dataset root
    :param demographics: dict from load_delaware_demographics(), or None
    :param n_valid: total number of patients in the valid split; must be
        divisible by 4 for a perfectly balanced (ptype × sex) grid
    :param n_test: total number of patients in the test split; same constraint
    :param seed: RNG seed for the patient shuffle within each sex group
    :return: dict {'train'/'valid'/'test': {audio_path: info_dict}}
    """
    import random

    data_folder = pathlib.Path(data_folder)
    demographics = demographics or {}
    splits = {"train": {}, "valid": {}, "test": {}}

    # n_valid // 4 and n_test // 4 patients per (class, sex) cell
    n_valid_per_cell = n_valid // 4
    n_test_per_cell = n_test // 4

    for subfolder, ptype in [("MCI", "Disease"), ("Control", "Control")]:
        sub_path = data_folder / subfolder
        if not sub_path.exists():
            print(f"Delaware subfolder not found, skipping: {sub_path}")
            continue

        class_demo = demographics.get(subfolder, {})

        # Group files by patient id -> list of (visit_str, wav_path)
        files_by_pid = {}
        for wav_path in sorted(glob.glob(str(sub_path / "*.wav"))):
            stem = pathlib.Path(wav_path).stem  # e.g. "12-3"
            try:
                pid, rec = stem.split("-", 1)
            except ValueError:
                print(f"Unexpected Delaware filename, skipping: {wav_path}")
                continue
            files_by_pid.setdefault(pid, []).append((_norm_visit(rec), wav_path))

        # Resolve demographics for each patient at their earliest recorded visit.
        def _resolve_demo(pid):
            recs = files_by_pid[pid]
            earliest_visit, _ = min(recs, key=lambda vr: _visit_sort_key(vr[0]))
            demo = class_demo.get((pid, earliest_visit))
            if demo is None and class_demo:
                pid_visits = sorted(
                    (v for (p, v) in class_demo if p == pid),
                    key=_visit_sort_key,
                )
                if pid_visits:
                    demo = class_demo[(pid, pid_visits[0])]
                    print(f"[delaware] {subfolder} pid={pid}: no demo row for "
                          f"visit={earliest_visit}, using visit={pid_visits[0]}")
                else:
                    print(f"[delaware] {subfolder} pid={pid}: no demographics found")
            return demo or {}, earliest_visit

        pids = sorted(files_by_pid.keys())

        # Collect (demo, earliest_visit, earliest_path) per pid.
        pid_info = {}
        for pid in pids:
            demo, earliest_visit = _resolve_demo(pid)
            earliest_path = min(
                files_by_pid[pid], key=lambda vr: _visit_sort_key(vr[0])
            )[1]
            pid_info[pid] = {
                "sex": demo.get("sex"),
                "earliest_visit": earliest_visit,
                "earliest_path": earliest_path,
                "demo": demo,
            }

        # Sex-stratified shuffle within this class.
        male_pids = [p for p in pids if pid_info[p]["sex"] == "M"]
        female_pids = [p for p in pids if pid_info[p]["sex"] == "F"]

        rng = random.Random(f"{seed}-{subfolder}")
        rng.shuffle(male_pids)
        rng.shuffle(female_pids)

        need = n_valid_per_cell + n_test_per_cell
        if len(male_pids) < need:
            print(f"[delaware] {subfolder}: only {len(male_pids)} male patients, "
                  f"need {need} for balanced split")
        if len(female_pids) < need:
            print(f"[delaware] {subfolder}: only {len(female_pids)} female patients, "
                  f"need {need} for balanced split")

        valid_pids = set(
            male_pids[:n_valid_per_cell] + female_pids[:n_valid_per_cell]
        )
        test_pids = set(
            male_pids[n_valid_per_cell:need] + female_pids[n_valid_per_cell:need]
        )

        for pid in pids:
            meta = pid_info[pid]
            info = {"ptype": ptype, "pid": f"{subfolder}-{pid}"}
            info.update(meta["demo"])
            info["visit"] = meta["earliest_visit"]  # recording-derived visit is authoritative

            if pid in test_pids:
                target = splits["test"]
            elif pid in valid_pids:
                target = splits["valid"]
            else:
                target = splits["train"]
            target[meta["earliest_path"]] = info

    return splits


# ---------------------------------------------------------------------------
# Combined JSON writer
# ---------------------------------------------------------------------------

# Unified schema. Any field absent for a given dataset is filled with None.
INFO_FIELDS = [
    "ptype", "sex", "age", "l1", "pid", "task", "lang",
    "moca", "visit", "test_date",
]


def _empty_info():
    return {k: None for k in INFO_FIELDS}


def create_combined_json(
    json_file,
    qpn_path_type_dict,
    pitt_ground_truth,
    delaware_path_type_dict,
    chunk_size,
    overlap=None,
    pitt_overlap=None,
    delaware_overlap=None,
):
    """
    Build a single JSON manifest combining QPN, Pitt, and Delaware entries.

    :param json_file: output path
    :param qpn_path_type_dict: dict {audio_path: trait_dict} for QPN
    :param pitt_ground_truth: pandas DataFrame for Pitt
    :param delaware_path_type_dict: dict {audio_path: trait_dict} for Delaware
    :param chunk_size: chunk size in seconds
    :param overlap: overlap arg for QPN entries (None -> hop = chunk/2)
    :param pitt_overlap: overlap arg for Pitt entries
    :param delaware_overlap: overlap arg for Delaware entries
    """
    json_dict = {}
    _add_qpn_entries(json_dict, qpn_path_type_dict, chunk_size, overlap)
    _add_pitt_entries(json_dict, pitt_ground_truth, chunk_size, pitt_overlap)
    _add_delaware_entries(
        json_dict, delaware_path_type_dict, chunk_size, delaware_overlap
    )

    with open(json_file, mode="w") as json_f:
        json_dict = convert_to_python(json_dict)
        json.dump(json_dict, json_f, indent=2)
    return json_dict

def _add_qpn_entries(json_dict, path_type_dict, chunk_size, overlap):
    hop_size = chunk_size / 2 if overlap is None else chunk_size - overlap

    for audiofile in path_type_dict.keys():
        # Skip 'l1' duplicate files
        if "l1" in audiofile:
            continue

        info_dict = _empty_info()
        info_dict.update(path_type_dict[audiofile])  # ptype, sex, age, l1

        audiopath = pathlib.Path(audiofile)
        items = audiopath.stem.split("_")
        info_dict.update({
            "pid": items[1],
            "task": items[2],
            "lang": items[-1],
        })

        # Task corrections
        if info_dict["task"] not in [
            "a1", "a2", "a3", "a4", "vowel_repeat", "dpt",
            "recall", "repeat", "hbd", "read",
        ]:
            info_dict["task"] = items[3]  # some batch 3 files
        if info_dict["task"] in ["a1", "a2", "a3", "a4"]:
            info_dict["task"] = "vowel_repeat"
        if info_dict["lang"] not in ["en", "fr"]:
            info_dict["lang"] = "other"

        info_dict["dataset"] = "qpn"

        audioinfo = sf.info(audiofile)
        duration = audioinfo.frames / audioinfo.samplerate

        # Use the file path itself as the unique base id
        base_id = audiofile

        max_start = max(duration - hop_size, 1)
        for i, start in enumerate(np.arange(0, max_start, hop_size)):
            chunk_duration = min(chunk_size, duration - i * hop_size)
            json_dict[f"{base_id}_{i}"] = {
                "wav": audiofile,
                "start": start,
                "duration": chunk_duration,
                "info_dict": info_dict,
            }


def _add_pitt_entries(json_dict, ground_truth, chunk_size, overlap):
    hop_size = chunk_size / 2 if overlap is None else chunk_size - overlap

    for row in ground_truth.to_dict(orient="records"):
        audioinfo = sf.info(row["path"])
        duration = audioinfo.frames / audioinfo.samplerate

        ptype = "Control" if row["dx"] == "Control" else "Disease"

        info_dict = _empty_info()
        info_dict.update({
            "ptype": ptype,
            "pid": str(row["id"]),
        })
        info_dict["dataset"] = "pitt"

        # Use the file path itself as the unique base id
        base_id = str(row["path"])

        max_start = max(duration - hop_size, 1)
        for i, start in enumerate(np.arange(0, max_start, hop_size)):
            chunk_duration = min(chunk_size, duration - i * hop_size)
            json_dict[f"{base_id}_{i}"] = {
                "wav": str(row["path"]),
                "start": start,
                "duration": chunk_duration,
                "info_dict": info_dict,
            }


def _add_delaware_entries(json_dict, path_type_dict, chunk_size, overlap):
    hop_size = chunk_size / 2 if overlap is None else chunk_size - overlap

    for audiofile, traits in path_type_dict.items():
        info_dict = _empty_info()
        info_dict.update(traits)  # ptype, pid, visit, age, sex, l1, moca, test_date
        info_dict["dataset"] = "delaware"

        audioinfo = sf.info(audiofile)
        duration = audioinfo.frames / audioinfo.samplerate

        # Use the file path itself as the unique base id
        base_id = audiofile

        max_start = max(duration - hop_size, 1)
        for i, start in enumerate(np.arange(0, max_start, hop_size)):
            chunk_duration = min(chunk_size, duration - i * hop_size)
            json_dict[f"{base_id}_{i}"] = {
                "wav": audiofile,
                "start": start,
                "duration": chunk_duration,
                "info_dict": info_dict,
            }
